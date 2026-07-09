from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DEFAULT_FONT = "Microsoft YaHei"
TITLE_COLOR = RGBColor(31, 92, 153)
TEXT_COLOR = RGBColor(31, 41, 51)
MUTED_COLOR = RGBColor(107, 114, 128)
HEADER_FILL = "1F5C99"


def set_run_font(run, size: int | None = None, bold: bool | None = None, color: RGBColor | None = None) -> None:
    run.font.name = DEFAULT_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), DEFAULT_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def configure_styles(doc: Document) -> None:
    for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet"]:
        style = doc.styles[style_name]
        style.font.name = DEFAULT_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), DEFAULT_FONT)

    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].font.color.rgb = TEXT_COLOR
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 1"].font.color.rgb = TITLE_COLOR
    doc.styles["Heading 2"].font.size = Pt(12)
    doc.styles["Heading 2"].font.bold = True
    doc.styles["Heading 2"].font.color.rgb = TEXT_COLOR

    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)


def add_paragraph(doc: Document, text: str, style: str | None = None, size: int | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.18
    run = paragraph.add_run(text)
    set_run_font(run, size=size)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(level=level)
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_run_font(run, size=15 if level == 1 else 12, bold=True, color=TITLE_COLOR if level == 1 else TEXT_COLOR)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_text(cell, value: Any, bold: bool = False, color: RGBColor | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("" if value is None else str(value))
    set_run_font(run, size=9, bold=bold, color=color or TEXT_COLOR)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_title_block(doc: Document, spec: dict[str, Any]) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(spec.get("title", "研究报告"))
    set_run_font(run, size=20, bold=True, color=TITLE_COLOR)

    subtitle = spec.get("subtitle")
    if subtitle:
        add_paragraph(doc, subtitle, size=11)

    meta_parts = []
    for key in ["author", "date", "source"]:
        if spec.get(key):
            meta_parts.append(str(spec[key]))
    if meta_parts:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        run = paragraph.add_run(" | ".join(meta_parts))
        set_run_font(run, size=8, color=MUTED_COLOR)


def add_summary(doc: Document, bullets: list[str]) -> None:
    if not bullets:
        return
    add_heading(doc, "核心观点", level=1)
    for item in bullets:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        set_run_font(run, size=10)


def add_table(doc: Document, table_spec: dict[str, Any]) -> None:
    title = table_spec.get("title")
    if title:
        add_paragraph(doc, title, size=9)

    columns = table_spec.get("columns") or []
    rows = table_spec.get("rows") or []
    if not columns:
        return

    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for idx, column in enumerate(columns):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, HEADER_FILL)
        set_cell_text(cell, column, bold=True, color=RGBColor(255, 255, 255))

    for row in rows:
        cells = table.add_row().cells
        for idx, column in enumerate(columns):
            value = row[idx] if isinstance(row, list) and idx < len(row) else row.get(column, "")
            set_cell_text(cells[idx], value)

    source = table_spec.get("source")
    if source:
        add_source_note(doc, source)


def add_image(doc: Document, image_spec: dict[str, Any], base_dir: Path) -> None:
    image_path = Path(image_spec.get("path", ""))
    if not image_path.is_absolute():
        image_path = base_dir / image_path
    if not image_path.exists():
        raise FileNotFoundError(image_path)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(float(image_spec.get("width_inches", 6.25))))

    caption = image_spec.get("caption")
    if caption:
        caption_para = doc.add_paragraph()
        caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_run = caption_para.add_run(caption)
        set_run_font(caption_run, size=8, color=MUTED_COLOR)

    source = image_spec.get("source")
    if source:
        add_source_note(doc, source)


def add_source_note(doc: Document, source: str) -> None:
    note = source if source.startswith("资料来源") else f"资料来源：{source}"
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(note)
    set_run_font(run, size=8, color=MUTED_COLOR)


def add_sections(doc: Document, spec: dict[str, Any], base_dir: Path) -> None:
    for section in spec.get("sections", []):
        add_heading(doc, section.get("heading", "正文分析"), level=1)
        for paragraph in section.get("paragraphs", []):
            add_paragraph(doc, paragraph)
        for table in section.get("tables", []):
            add_table(doc, table)
        for image in section.get("images", []):
            add_image(doc, image, base_dir)


def add_risk_warnings(doc: Document, warnings: list[str]) -> None:
    if not warnings:
        return
    add_heading(doc, "风险提示与局限说明", level=1)
    for item in warnings:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_run_font(run, size=10)


def build_docx(spec_path: Path, output_path: Path) -> Path:
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    doc = Document()
    configure_styles(doc)
    add_title_block(doc, spec)
    add_summary(doc, spec.get("summary", []))
    add_sections(doc, spec, spec_path.parent)
    add_risk_warnings(doc, spec.get("risk_warnings", []))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build a Chinese research-report-style DOCX from a JSON spec.")
    parser.add_argument("--spec", required=True, help="Report JSON spec path.")
    parser.add_argument("--output", required=True, help="Output DOCX path.")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    path = build_docx(Path(args.spec), Path(args.output))
    print(path)
